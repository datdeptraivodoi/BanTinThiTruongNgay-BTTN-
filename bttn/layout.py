"""Word layout primitives for the approved Times New Roman 11 template."""
from copy import deepcopy

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Times New Roman"
SIZE = 11
NAVY = "245794"
FILL = "4F81BD"


def style_run(run, font=FONT, size=SIZE):
    run.font.name = font
    run.font.size = Pt(size)
    properties = run._r.get_or_add_rPr()
    fonts = properties.get_or_add_rFonts()
    for name in list(fonts.attrib):
        if name.endswith("Theme"):
            del fonts.attrib[name]
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + name), font)
    val_sz = str(int(size * 2))
    properties.get_or_add_sz().set(qn("w:val"), val_sz)
    size_cs = properties.find(qn("w:szCs"))
    if size_cs is None:
        size_cs = OxmlElement("w:szCs")
        properties.append(size_cs)
    size_cs.set(qn("w:val"), val_sz)


def compact_paragraph(p):
    fmt = p.paragraph_format
    fmt.space_before = fmt.space_after = Pt(0)
    fmt.line_spacing = 1
    fmt.left_indent = fmt.right_indent = fmt.first_line_indent = Pt(0)
    fmt.keep_with_next = False
    fmt.widow_control = False
    for run in p.runs:
        sz_node = run._r.find(qn("w:rPr"))
        val = sz_node.find(qn("w:sz")).get(qn("w:val")) if sz_node is not None and sz_node.find(qn("w:sz")) is not None else None
        r_fonts = sz_node.find(qn("w:rFonts")) if sz_node is not None else None
        font_name = r_fonts.get(qn("w:ascii")) if r_fonts is not None else run.font.name
        if font_name == "Verdana":
            size = 18 if val == "36" else 12
            style_run(run, font="Verdana", size=size)
        elif val == "20":
            style_run(run, size=10)
        elif val == "18":
            style_run(run, size=9)
        elif val == "16":
            style_run(run, size=8)
        else:
            style_run(run, size=SIZE)
    if not p.text and not p._p.findall(".//" + qn("w:drawing")):
        tr = p._p
        while tr is not None and tr.tag != qn("w:tr"):
            tr = tr.getparent()
        if tr is not None and any(shd.get(qn("w:fill")) == "4F81BD" for shd in tr.iter(qn("w:shd"))):
            fmt.line_spacing = Pt(12)
            fmt.space_before = fmt.space_after = Pt(0)
        else:
            # Required empty paragraphs after nested tables are not visible text.
            fmt.line_spacing = Pt(1)
            fmt.space_before = fmt.space_after = Pt(0)


def normalize_typography(doc):
    from docx.text.paragraph import Paragraph

    for style in doc.styles:
        if style.type in (1, 2):
            style.font.name = FONT
            style.font.size = Pt(SIZE)
    for part in [doc.part]:
        for node in part.element.iter(qn("w:p")):
            compact_paragraph(Paragraph(node, part))
    for node in list(doc.element.iter(qn("w:trHeight"))):
        tr = node.getparent()
        if tr is not None and any(shd.get(qn("w:fill")) == "4F81BD" for shd in tr.iter(qn("w:shd"))):
            continue
        if tr is not None:
            tr.remove(node)
    for node in doc.element.iter(qn("w:tblHeader")):
        node.getparent().remove(node)
    for node in doc.element.iter(qn("w:tcMar")):
        for edge in node:
            edge.set(qn("w:w"), "10")
    for s in doc.sections:
        for ref in list(s._sectPr):
            if ref.tag in {qn("w:headerReference"), qn("w:footerReference")}:
                s._sectPr.remove(ref)


def styled_table(cell, rows, ratios, blue_cells=(), red_cells=()):
    width = int((cell.width or Inches(3.5)) - Inches(.08))
    table = cell.add_table(rows=len(rows), cols=len(ratios))
    table.autofit = False
    for column, ratio in zip(table.columns, ratios):
        column.width = int(width * ratio)
    for i, values in enumerate(rows):
        table.rows[i]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for j, text in enumerate(values):
            c = table.cell(i, j)
            c.width = int(width * ratios[j])
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_together = True
            r = p.add_run(str(text))
            style_run(r)
            r.bold = True
            r.font.color.rgb = RGBColor.from_string(
                "FFFFFF" if (i, j) in blue_cells else "C00000" if (i, j) in red_cells else NAVY)
            if (i, j) in blue_cells:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), FILL)
                c._tc.get_or_add_tcPr().append(shade)
            margins = OxmlElement("w:tcMar")
            for name in ("top", "bottom", "left", "right"):
                edge = OxmlElement("w:" + name)
                edge.set(qn("w:w"), "10")
                edge.set(qn("w:type"), "dxa")
                margins.append(edge)
            c._tc.get_or_add_tcPr().append(margins)
    table._tbl.tblPr.find(qn("w:tblW")).set(qn("w:w"), str(round(width / 635)))
    table._tbl.tblPr.find(qn("w:tblW")).set(qn("w:type"), "dxa")
    return table


def split_pages(doc, table):
    """Explicit page boundaries; separator rows cannot drift across pages."""
    parent = table._tbl.getparent()
    position = parent.index(table._tbl)
    for indexes in [range(0, 6), range(6, 10), range(10, 14)]:
        if indexes.start:
            p = OxmlElement("w:p")
            pr = OxmlElement("w:pPr")
            pr.append(OxmlElement("w:pageBreakBefore"))
            spacing = OxmlElement("w:spacing")
            spacing.set(qn("w:before"), "0")
            spacing.set(qn("w:after"), "0")
            spacing.set(qn("w:line"), "1")
            spacing.set(qn("w:lineRule"), "exact")
            pr.append(spacing)
            p.append(pr)
            r = OxmlElement("w:r")
            br = OxmlElement("w:br")
            br.set(qn("w:type"), "page")
            r.append(br)
            p.append(r)
            parent.insert(position, p)
            position += 1
        tbl = deepcopy(table._tbl)
        for child in list(tbl):
            if child.tag == qn("w:tr"):
                tbl.remove(child)
        for index in indexes:
            tbl.append(deepcopy(table.rows[index]._tr))
        parent.insert(position, tbl)
        position += 1
    parent.remove(table._tbl)
