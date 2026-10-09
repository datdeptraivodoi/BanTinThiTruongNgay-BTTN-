import json
from datetime import timedelta
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from bttn.models import ReportContent, Snapshot, parse_as_of
from bttn.rendering import render, render_sbv_block
from bttn.trader_quotes import (
    QuoteFile,
    apply_trader_quotes,
    collect_trader_quotes,
    validate_trader_observations,
)
from bttn.validation import count_words, resolve, validate_content

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def snapshot():
    return Snapshot.model_validate_json((ROOT / 'tests/fixtures/snapshot.json').read_text(encoding='utf-8'))


@pytest.fixture
def content():
    return ReportContent.model_validate_json((ROOT / 'tests/fixtures/content.json').read_text(encoding='utf-8'))


@pytest.mark.parametrize('index,count,valid', [(0,34,False), (0,35,True), (0,45,True),
                                              (0,46,False), (1,34,False), (1,35,True), (2,12,True)])
def test_highlight_length_after_substitution(snapshot, content, index, count, valid):
    content.highlights[index].paragraphs = ['{{GOLD}} ' + 'tin ' * (count - 1)]
    failures = [i for i in validate_content(content, snapshot)
                if i.code == 'WORD_COUNT' and i.message.startswith(f'highlight_{index}:')]
    assert bool(failures) == (not valid)
    assert count_words(resolve(' '.join(content.highlights[index].paragraphs), snapshot)) == count


def test_word_count_excludes_display_numbering():
    assert count_words('1. Đây là tin') == 3
    assert count_words('2) Giá vàng 2.700 USD') == 4


def quote_file():
    return QuoteFile.model_validate_json((ROOT / 'config/trader_quotes.example.json').read_text(encoding='utf-8'))


def remove_trader_data(snapshot):
    snapshot.observations = {k: v for k, v in snapshot.observations.items()
                             if not k.startswith(('INTERBANK_', 'ALM_SWAP_'))}


def test_future_and_stale_quotes_are_not_current(snapshot):
    remove_trader_data(snapshot)
    data = quote_file()
    data.quotes[-1].quoted_at = snapshot.as_of + timedelta(minutes=1)
    apply_trader_quotes(snapshot, data)
    assert 'INTERBANK_BID' not in snapshot.observations
    assert 'ALM_SWAP_ON_BID' not in snapshot.observations
    assert snapshot.observations['INTERBANK_BID_PREV'].trading_date.isoformat() == '2026-09-23'


def test_monday_uses_friday_quote_not_weekend(snapshot):
    remove_trader_data(snapshot)
    snapshot.as_of = parse_as_of('2026-09-28T12:00:00+07:00')
    data = quote_file()
    data.quotes[0].quoted_at = parse_as_of('2026-09-25T11:55:00+07:00')
    data.quotes[-1].quoted_at = parse_as_of('2026-09-27T11:55:00+07:00')
    apply_trader_quotes(snapshot, data)
    assert snapshot.observations['INTERBANK_BID_PREV'].trading_date.isoformat() == '2026-09-25'
    assert 'INTERBANK_BID' not in snapshot.observations


def test_test_quote_file_is_refused_in_production(snapshot):
    snapshot.purpose = 'live'
    with pytest.raises(ValueError, match='Fixture'):
        apply_trader_quotes(snapshot, quote_file())


@pytest.mark.parametrize('change', ['fx_order', 'fx_half', 'swap_order', 'duplicate', 'no_timezone'])
def test_invalid_trader_quotes_rejected(change):
    data = json.loads((ROOT / 'config/trader_quotes.example.json').read_text(encoding='utf-8'))
    quote = data['quotes'][-1]
    if change == 'fx_order':
        quote['fx_bid'] = '27000'
    elif change == 'fx_half':
        del quote['fx_ask']
    elif change == 'swap_order':
        quote['swaps'][0]['bid'] = '10'
    elif change == 'duplicate':
        quote['swaps'].append(quote['swaps'][0])
    else:
        quote['quoted_at'] = '2026-09-24T11:55:00'
    with pytest.raises(ValueError):
        QuoteFile.model_validate(data)


def test_no_quote_configuration_does_not_fabricate(snapshot, monkeypatch):
    remove_trader_data(snapshot)
    monkeypatch.delenv('TRADER_QUOTES_PATH', raising=False)
    collect_trader_quotes(None, snapshot)
    assert not any(k.startswith('ALM_SWAP_') for k in snapshot.observations)
    assert snapshot.issues[-1].code == 'TRADER_QUOTES_MISSING'


@pytest.mark.parametrize('error', ['stale', 'future', 'inverted', 'half_pair'])
def test_saved_snapshot_trader_quotes_are_revalidated(snapshot, error):
    bid = snapshot.observations['INTERBANK_BID']
    if error == 'stale':
        bid.trading_date -= timedelta(days=1)
    elif error == 'future':
        snapshot.sources[bid.source_id].published_at = snapshot.as_of + timedelta(minutes=1)
    elif error == 'inverted':
        bid.value = snapshot.observations['INTERBANK_ASK'].value + 1
    else:
        snapshot.observations.pop('INTERBANK_ASK')
    assert validate_trader_observations(snapshot)


def test_sbv_layout_and_no_mb_substitution(snapshot):
    doc = Document()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    render_sbv_block(cell, snapshot)
    fixing, interbank = cell.tables
    assert (len(fixing.rows), len(fixing.columns)) == (4, 3)
    assert fixing.cell(2, 0).text == fixing.cell(3, 0).text == ''
    assert fixing.cell(1, 0)._tc.tcPr.find(qn('w:shd')).get(qn('w:fill')) == '4F81BD'
    assert interbank.cell(1, 1).text == '25.800/26.000'
    assert interbank.cell(1, 1).paragraphs[0].runs[0].font.color.rgb.__str__() == 'C00000'
    snapshot.observations.pop('MB_BUY', None)
    snapshot.observations.pop('MB_SELL', None)
    remove_trader_data(snapshot)
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    render_sbv_block(cell, snapshot)
    assert cell.tables[1].cell(1, 1).text == '—'


def test_renderer_uses_normal_prose_and_applies_font(snapshot, content, tmp_path):
    # Regression: ordinary prose used to be discarded unless it contained a test marker.
    content.highlights[0].paragraphs = ['Tin biên tập đã duyệt dùng để kiểm tra nội dung xuất ra.']
    content.interbank.paragraphs = ['Nội dung liên ngân hàng đã được biên tập.']
    output = tmp_path / 'preview.docx'
    render(snapshot, content, ROOT / 'template.docx', output)
    doc = Document(output)
    assert len(doc.tables) == 3
    xml_text = '\n'.join(node.text or '' for node in doc.element.iter(qn('w:t')))
    assert content.highlights[0].paragraphs[0] in xml_text
    assert content.interbank.paragraphs[0] in xml_text
    assert 'lãi suất ON nhiều khả năng đi ngang' in xml_text
    assert 'Sử dụng sản phẩm' in xml_text
    assert 'MBBank · chuyển khoản' not in xml_text
    for run in doc.element.iter(qn('w:r')):
        if not list(run.iter(qn('w:t'))):
            continue
        props = run.find(qn('w:rPr'))
        assert props.find(qn('w:sz')).get(qn('w:val')) == '22'
        assert props.find(qn('w:rFonts')).get(qn('w:ascii')) == 'Times New Roman'
