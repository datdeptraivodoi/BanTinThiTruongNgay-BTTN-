"""Phản hồi review 15/10: giá MBBank, bảng dự báo, lịch sự kiện, Term SOFR, VIRA, bố cục theo mẫu."""
import json
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from bttn import rendering
from bttn.models import Snapshot
from bttn.rendering import (
    TEMPLATE_GRID,
    TEMPLATE_SPANS,
    calendar_events,
    forecast_rows,
    render_forecast_table,
    term_sofr_rows,
)
from bttn.sources import _webgia_hex, parse_vira_sections, parse_webgia_mb

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def snap():
    return Snapshot.model_validate_json((ROOT / "tests/fixtures/snapshot.json").read_text(encoding="utf-8"))


def test_webgia_hex_roundtrip():
    assert _webgia_hex("3235XQ2eZ38W30") == "25.80"  # chữ hoa là nhiễu, bị loại


def test_parse_webgia_mb_requires_today_and_reads_usd_row():
    html = ("<p>Tỷ giá MBBank cập nhật lúc 09:15 15/10/2026</p><table>"
            "<tr><th>Mã</th><th>Tên</th><th>Mua TM</th><th>Mua CK</th><th>Bán</th></tr>"
            "<tr><td>USD (50-100)</td><td>Đô la Mỹ</td><td>25.900</td><td>25.940</td><td>26.290</td></tr></table>")
    assert parse_webgia_mb(html, date(2026, 10, 15)) == (25940, 26290)
    with pytest.raises(ValueError):
        parse_webgia_mb(html, date(2026, 10, 16))


def test_forecast_rows_follow_config_and_go_blank_when_stale(snap, monkeypatch):
    cfg = {"updated": snap.as_of.date().isoformat(), "max_age_days": 45,
           "rows": [{"name": "USD-VND", "1m": "26,150", "6m": "26,400", "note": "Ghi chú"}]}
    monkeypatch.setattr(rendering, "_config_json", lambda name: cfg)
    assert forecast_rows(snap) == [("USD-VND", "26,150", "26,400", "Ghi chú")]
    cfg["updated"] = "2020-01-01"
    assert forecast_rows(snap)[0][1:3] == ("—", "—")
    assert snap.issues[-1].code == "FORECAST_STALE"


def test_forecast_table_matches_template_header(snap):
    cell = Document().add_table(rows=1, cols=1).cell(0, 0)
    render_forecast_table(cell, snap)
    assert cell.paragraphs[0].text == "Dự báo các chỉ số chính" or "Dự báo các chỉ số chính" in [p.text for p in cell.paragraphs]
    table = cell.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["", "1 tháng", "6 tháng", "Ghi chú"]
    assert [r.cells[0].text for r in table.rows[1:]] == ["USD-VND", "EUR-USD", "USD-JPY", "LiborUSD 6M"]


def test_calendar_uses_only_the_current_week(snap, monkeypatch):
    cfg = {"week_start": "2026-09-21", "week_end": "2026-09-25", "events": [
        {"date": "2026-09-22", "event": "A", "country": "Mỹ", "high_impact": True},
        {"date": "2026-09-22", "event": "B", "country": "Mỹ"},
        {"date": "2026-09-23", "event": "C", "country": "Nhật"}]}
    monkeypatch.setattr(rendering, "_config_json", lambda name: cfg)
    week, events = calendar_events(snap)
    assert [e[0] for e in events] == ["22.09.2026", "", "23.09.2026"]
    cfg["week_start"], cfg["week_end"] = "2026-10-12", "2026-10-16"
    assert calendar_events(snap) == (None, [])


def test_term_sofr_only_when_recent(snap, monkeypatch):
    cfg = {"date": snap.as_of.date().isoformat(), "1M": 4.1, "3M": 4.05, "6M": None}
    monkeypatch.setattr(rendering, "_config_json", lambda name: cfg)
    rows = term_sofr_rows(snap)
    assert rows[0][1] == "4.10000%" and rows[2][1] == "—"
    cfg["date"] = "2026-09-01"
    assert all(r[1] == "—" for r in term_sofr_rows(snap))


def test_config_files_are_valid_json():
    for name in ("forecasts.json", "economic_calendar.json", "term_sofr.json", "policy_rates.json"):
        json.loads((ROOT / "config" / name).read_text(encoding="utf-8"))


def test_vira_sections_split_fx_and_money_market():
    text = ("Thị trường ngoại tệ (menu)\nBản tin\nThị trường ngoại tệ:\nTỷ giá trên thị trường tự do ở mức 26.020/26.120.\n"
            "Thị trường tiền tệ LNH:\nLãi suất qua đêm giảm.\nNghiệp vụ thị trường mở:\nNHNN hút ròng 764,91 tỷ đồng.\n"
            "Thị trường chứng khoán: VN-Index giảm.")
    out = parse_vira_sections(text)
    assert out["fx"] == "Tỷ giá trên thị trường tự do ở mức 26.020/26.120."
    assert out["lnh"] == "Lãi suất qua đêm giảm.\n\nNHNN hút ròng 764,91 tỷ đồng."


def test_template_geometry_matches_uploaded_template():
    assert sum(TEMPLATE_GRID) == 15877
    for row, spans in TEMPLATE_SPANS.items():
        assert sum(spans) + rendering.TEMPLATE_GRID_BEFORE.get(row, 0) == len(TEMPLATE_GRID)
    def widths(row):
        start, out = rendering.TEMPLATE_GRID_BEFORE.get(row, 0), []
        for span in TEMPLATE_SPANS[row]:
            out.append(sum(TEMPLATE_GRID[start:start + span]))
            start += span
        return out
    # Đúng độ rộng ô của file mẫu BT.TT.2026.10.19.docx
    assert widths(3) == [5157, 4610, 6076]
    assert widths(7) == [5297, 5221, 5325]
    assert widths(11) == [5323, 5323, 5231]


def test_rendered_word_keeps_template_columns(tmp_path):
    from bttn.pipeline import main
    out = tmp_path / "out"
    assert main(["--create-only", "--snapshot", str(ROOT / "tests/fixtures/snapshot.json"),
                 "--content", str(ROOT / "tests/fixtures/content.json"), "--output-dir", str(out)]) == 0
    doc = Document(next(out.rglob("BTTN-*.docx")))
    page1 = doc.tables[0]
    grid = [int(g.get(qn("w:w"))) for g in page1._tbl.tblGrid.findall(qn("w:gridCol"))]
    assert grid == TEMPLATE_GRID
    texts = "".join(t.text or "" for t in doc.element.iter(qn("w:t")))
    assert "Dự báo các chỉ số chính" in texts
    assert "Tỷ giá USD-VND ngân hàng" in texts


def test_eur_usd_has_europe_paragraph_then_us_paragraph(snap):
    from datetime import timedelta
    from types import SimpleNamespace

    from bttn.editorial import foreign_sections
    from bttn.models import Source
    from bttn.validation import count_words, resolve

    def te(topic_path, i, hours, text):
        sid = f"news_te_{topic_path}_{i}"
        at = snap.as_of - timedelta(hours=hours)
        snap.sources[sid] = Source(id=sid, url=f"https://tradingeconomics.com/{topic_path}/news/{i}",
                                   published_at=at, retrieved_at=at, text="Title\n\nBody", kind="news",
                                   text_scope="article")
        return sid, SimpleNamespace(text=text)

    us = ("Lạm phát CPI của Mỹ trong tháng 9 tăng 0,6% so với tháng trước, mức tăng mạnh nhất trong nhiều tháng qua. "
          "Giá xăng tăng kỷ lục là nguyên nhân chính khiến chỉ số giá tiêu dùng của Mỹ đi lên. "
          "Lạm phát lõi của Mỹ tăng 2,6% so với cùng kỳ năm trước, thấp hơn mức của tháng 8. "
          "Thị trường đánh giá khả năng Fed tăng lãi suất trong cuộc họp cuối tháng 10 ở mức cao hơn. ")
    eu = ("Lạm phát khu vực đồng euro trong tháng 9 tăng lên 2,4% so với cùng kỳ năm trước theo số liệu sơ bộ. "
          "Sản lượng công nghiệp của Đức giảm 1,1% so với tháng trước, kém hơn mức kỳ vọng của thị trường. "
          "Chỉ số niềm tin kinh tế ZEW của Đức giảm xuống mức thấp nhất trong sáu tháng gần đây. "
          "Doanh số bán lẻ của khu vực đồng euro tăng 0,3% so với tháng trước nhờ chi tiêu cho dịch vụ. "
          "Tỷ lệ thất nghiệp của khu vực đồng euro ổn định ở mức 6,3% trong tháng 8 vừa qua. ")
    records = dict([te("united-states", 1, 3, us), te("euro-area", 2, 5, eu), te("germany", 3, 20, eu)])
    result, audit = foreign_sections(snap, records)
    section = result["eur_usd"]
    assert len(section.paragraphs) == 2
    first, second = section.paragraphs
    assert first.startswith("Trong phiên giao dịch hôm qua, tính đến ngày")
    assert "khu vực đồng euro" in first or "Đức" in first
    assert second.startswith("- Về phía kinh tế Mỹ,") and "của Mỹ" in second
    assert 185 <= count_words(resolve(first + " " + second, snap, safe=True)) <= 200, audit["eur_usd"]


def test_calendar_highlights_the_bulletin_date(snap, monkeypatch):
    """Góp ý 19.10: sự kiện của ngày viết bản tin in đậm xanh đậm, ngày khác chữ thường xanh nhạt."""
    from datetime import timedelta
    day = snap.as_of.date()
    cfg = {"week_start": (day - timedelta(days=day.weekday())).isoformat(),
           "week_end": (day - timedelta(days=day.weekday()) + timedelta(days=4)).isoformat(),
           "events": [{"date": (day - timedelta(days=1)).isoformat(), "event": "Hôm qua", "country": "Mỹ"},
                      {"date": day.isoformat(), "event": "Hôm nay 1", "country": "Việt Nam"},
                      {"date": day.isoformat(), "event": "Hôm nay 2", "country": "Trung Quốc"},
                      {"date": (day + timedelta(days=1)).isoformat(), "event": "Ngày mai", "country": "Anh Quốc"}]}
    monkeypatch.setattr(rendering, "_config_json", lambda name: cfg)
    _, events = calendar_events(snap)
    assert [(e[1], e[3]) for e in events] == [("Hôm qua", False), ("Hôm nay 1", True), ("Hôm nay 2", True), ("Ngày mai", False)]
    doc = Document()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    rendering.render_calendar_events_table(cell, snap)
    rows = cell.tables[0].rows[1:]
    runs = [r.cells[1].paragraphs[0].runs[0] for r in rows]
    assert [bool(r.bold) for r in runs] == [False, True, True, False]
    assert [str(r.font.color.rgb) for r in runs] == ["8DB3E2", "4F81BD", "4F81BD", "8DB3E2"]


def test_calendar_falls_back_to_automatic_week(snap, monkeypatch):
    from datetime import timedelta
    monkeypatch.setattr(rendering, "_config_json", lambda name: {"week_start": "2020-01-06", "week_end": "2020-01-10"})
    day = snap.as_of.date()
    week, events = calendar_events(snap, auto=lambda monday: [(day, "CPI", "Mỹ"), (day + timedelta(days=1), "GDP", "Anh Quốc")])
    assert week[0].weekday() == 0 and events[0][3] is True and events[1][3] is False
    assert any(i.code == "CALENDAR_AUTO" for i in snap.issues)


def test_forecast_table_colors_and_numbers_follow_sample(snap, monkeypatch):
    """Góp ý 19.10: số trong bảng dự báo lấy theo file mẫu; tên đậm đen, số và ghi chú chữ thường đen."""
    cfg = json.loads((ROOT / "config/forecasts.json").read_text(encoding="utf-8"))
    cfg["updated"] = snap.as_of.date().isoformat()     # số giữ nguyên, chỉ đặt ngày trùng snapshot test
    monkeypatch.setattr(rendering, "_config_json", lambda name: cfg)
    doc = Document()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    rendering.render_forecast_table(cell, snap)
    table = cell.tables[0]
    assert [c.text for c in table.rows[1].cells] == ["USD-VND", "25,900", "26,150", "Thanh khoản ngoại tệ vẫn ở mức tốt, thâm hụt thu hẹp"]
    label, value, note = (table.rows[2].cells[i].paragraphs[0].runs[0] for i in (0, 1, 3))
    assert label.bold and not value.bold and not note.bold
    assert {str(r.font.color.rgb) for r in (label, value, note)} == {"000000"}
    assert table._tbl.tblPr.find(qn("w:tblBorders")) is not None
