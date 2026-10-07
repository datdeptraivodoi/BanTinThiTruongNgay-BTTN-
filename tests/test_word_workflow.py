import json
from pathlib import Path
from unittest.mock import Mock

import pytest
import yaml
from docx import Document
from docx.shared import Pt

from bttn import delivery, pipeline, rendering
from bttn.models import Snapshot

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "tests/fixtures"


def write_snapshot(tmp_path, *, live=False):
    snapshot = Snapshot.model_validate_json((FIXTURES / "snapshot.json").read_text(encoding="utf-8"))
    if live:
        snapshot.purpose = "live"
    path = tmp_path / "snapshot.json"
    path.write_text(snapshot.model_dump_json(), encoding="utf-8")
    return snapshot, path


def run_args(tmp_path, snapshot_path):
    return ["--snapshot", str(snapshot_path), "--content", str(FIXTURES / "content.json"),
            "--output-dir", str(tmp_path / "output"), "--state-dir", str(tmp_path / "state")]


def manifest(tmp_path):
    return json.loads(next((tmp_path / "output").glob("*/manifest.json")).read_text(encoding="utf-8"))


def test_workflow_offers_two_actions_without_smtp_sample_or_recipient_override():
    text = (ROOT / ".github/workflows/market_report.yml").read_text(encoding="utf-8")
    workflow = yaml.safe_load(text)
    events = workflow.get("on") or workflow[True]
    inputs = events["workflow_dispatch"]["inputs"]
    assert set(inputs) == {"action"}
    action = inputs["action"]
    assert action["type"] == "choice"
    assert action["default"] == "1. Gửi file Word đến các email"
    assert action["options"] == ["1. Gửi file Word đến các email", "2. Chỉ tạo file Word"]
    assert "--test-smtp" not in text and "--test-recipient" not in text
    assert "Regression tests" not in text and "libreoffice" not in text.lower()


def test_create_only_produces_one_word_without_pdf_converter_or_email(tmp_path, monkeypatch):
    _, path = write_snapshot(tmp_path)
    convert = Mock(side_effect=AssertionError("PDF conversion must remain disabled"))
    send = Mock(side_effect=AssertionError("Create-only must not send email"))
    monkeypatch.setattr(rendering, "convert_and_validate", convert)
    monkeypatch.setattr(delivery, "send_report", send)
    assert pipeline.main(["--create-only", *run_args(tmp_path, path)]) == 0
    result = manifest(tmp_path)
    assert result["status"] == "word_created" and result["send_requested"] is False
    assert result["pdf_generated"] is False
    assert len(result["artifacts"]) == 1 and next(iter(result["artifacts"])).endswith(".docx")
    assert list((tmp_path / "output").glob("*/*.docx")) and not list(tmp_path.rglob("*.pdf"))
    convert.assert_not_called()
    send.assert_not_called()


def test_send_uses_official_recipients_and_only_the_current_word(tmp_path, monkeypatch):
    _, path = write_snapshot(tmp_path, live=True)
    send = Mock()
    convert = Mock(side_effect=AssertionError("No hidden PDF generation"))
    monkeypatch.setenv("RECIPIENTS", "one@example.com, two@example.com")
    monkeypatch.setattr(delivery, "send_report", send)
    monkeypatch.setattr(rendering, "convert_and_validate", convert)
    assert pipeline.main(["--send", *run_args(tmp_path, path)]) == 0
    attachments = send.call_args.args[1]
    assert len(attachments) == 1 and attachments[0].suffix == ".docx"
    assert attachments[0].is_file()
    assert send.call_args.args[3] == ["one@example.com", "two@example.com"]
    assert manifest(tmp_path)["status"] == "sent"
    assert manifest(tmp_path)["recipients"] == ["one@example.com", "two@example.com"]
    assert not list(tmp_path.rglob("*.pdf"))
    convert.assert_not_called()


def test_bad_word_typography_blocks_delivery(tmp_path, monkeypatch):
    _, path = write_snapshot(tmp_path, live=True)
    original = rendering.render

    def bad_render(*args, **kwargs):
        original(*args, **kwargs)
        output = args[3]
        doc = Document(output)
        run = next(r for p in doc.tables[0].cell(0, 0).paragraphs for r in p.runs if r.text.strip())
        run.font.size = Pt(13)
        doc.save(output)

    send = Mock()
    monkeypatch.setattr(rendering, "render", bad_render)
    monkeypatch.setattr(delivery, "send_report", send)
    assert pipeline.main(["--send", *run_args(tmp_path, path)]) == 1
    send.assert_not_called()
    assert "11 pt" in manifest(tmp_path)["error_message"]


def test_smtp_attaches_one_word_without_test_subject_or_pdf(tmp_path, monkeypatch):
    snapshot, _ = write_snapshot(tmp_path, live=True)
    monkeypatch.setenv("SENDER_EMAIL", "sender@example.com")
    monkeypatch.setenv("SENDER_PASSWORD", "test-only-password")
    server = Mock()
    server.send_message.return_value = {}
    context = Mock()
    context.__enter__ = Mock(return_value=server)
    context.__exit__ = Mock(return_value=False)
    monkeypatch.setattr(delivery.smtplib, "SMTP_SSL", Mock(return_value=context))
    word = tmp_path / "report.docx"
    word.write_bytes(b"test-only attachment")
    ledger = delivery.send_report(snapshot, [word], tmp_path / "state", ["one@example.com"])
    message = server.send_message.call_args.args[0]
    attachments = list(message.iter_attachments())
    assert len(attachments) == 1 and attachments[0].get_filename() == "report.docx"
    assert attachments[0].get_content_type() == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert "TEST" not in message["Subject"] and message["To"] == "one@example.com"
    assert json.loads(ledger.read_text(encoding="utf-8"))["status"] == "sent"


def test_removed_smtp_sample_flag_is_rejected():
    with pytest.raises(SystemExit) as exc:
        pipeline.main(["--test-smtp"])
    assert exc.value.code == 2
