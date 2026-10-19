"""Góp ý bản tin 19.10 (tối): file 5 trang, tin nổi bật, định dạng trang 2–3 theo bản mẫu, mục vàng một câu."""
from datetime import timedelta
from decimal import Decimal
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from bttn import editorial
from bttn.models import Snapshot, Source
from bttn.validation import resolve

ROOT = Path(__file__).resolve().parents[1]


def _render(tmp_path):
    from bttn.pipeline import main
    out = tmp_path / "out"
    assert main(["--create-only", "--snapshot", str(ROOT / "tests/fixtures/snapshot.json"),
                 "--content", str(ROOT / "tests/fixtures/content.json"), "--output-dir", str(out)]) == 0
    return Document(next(out.rglob("BTTN-*.docx")))


def _para(doc, prefix):
    from docx.text.paragraph import Paragraph
    for node in doc.element.body.iter(qn("w:p")):
        p = Paragraph(node, None)
        if p.text.strip().startswith(prefix):
            return p
    raise AssertionError(prefix)


def test_two_page_breaks_and_highlights_not_spaced(tmp_path):
    doc = _render(tmp_path)
    # Đúng hai ngắt trang (trước trang 2 và trang 3)
    breaks = [br for br in doc.element.body.iter(qn("w:br")) if br.get(qn("w:type")) == "page"]
    breaks += [pb for pb in doc.element.body.iter(qn("w:pageBreakBefore")) if pb.get(qn("w:val")) not in ("0", "false")]
    assert len(breaks) == 2
    news = [p for p in doc.tables[0].rows[3].cells[0].paragraphs if p.text.strip()]
    assert all(p.paragraph_format.space_after == 0 for p in news)


def test_titles_and_contact_follow_sample(tmp_path):
    doc = _render(tmp_path)
    for prefix in ("Sự kiện quan trọng", "Bảng giá hàng hóa ngày", "Thị trường tiền tệ quốc tế"):
        run = _para(doc, prefix).runs[0]
        assert run.font.name == "Verdana" and run.font.size.pt == 10 and run.font.bold, prefix
    for prefix in ("Thông tin liên hệ", "Miễn trừ trách nhiệm"):
        run = _para(doc, prefix).runs[0]
        assert run.font.size.pt == 14 and run.font.bold and str(run.font.color.rgb) == "1F497D", prefix
    assert _para(doc, "Phụ trách Bộ phận Phân tích").runs[0].font.bold
    # Miễn trừ trách nhiệm: dòng trống giữa các đoạn như bản mẫu
    cell = doc.tables[2].rows[3].cells[0]
    texts = [p.text.strip() for p in cell.paragraphs]
    start = texts.index("Miễn trừ trách nhiệm")
    body = texts[start:]
    assert all(body[i] == "" for i in range(1, len(body) - 1, 2)) and body[-1]
    # Lịch sự kiện: dòng sự kiện giãn 1.2
    cal = doc.tables[1].rows[1].cells[0].tables[0]
    assert cal.rows[1].cells[1].paragraphs[0].paragraph_format.line_spacing == 1.2


def test_highlight_strips_glued_headline_and_skips_connectors():
    title = "Việt Nam xuất siêu sang EU 36 tỷ USD"
    glued = title + " Thương mại Việt Nam – EU tăng trưởng tích cực trong 9 tháng năm 2026."
    assert editorial.strip_heading(glued, title).startswith("Thương mại Việt Nam – EU")
    s = Snapshot.model_validate_json((ROOT / "tests/fixtures/snapshot.json").read_text(encoding="utf-8"))
    body = (title + "\n\n" + glued + " Tuy nhiên, WB cũng phải đối mặt với sự sụt giảm mạnh trong nguồn viện trợ "
            "phát triển chính thức và xuất khẩu.")
    s.sources["news_x"] = Source(id="news_x", url="https://vneconomy.vn/x.htm", published_at=s.as_of - timedelta(hours=2),
                                 retrieved_at=s.as_of, text=body, kind="news", text_scope="article")
    cands, rejected = editorial.news_candidates(s, {}, "macro")
    texts = [c.text for c in cands if c.source_id == "news_x"]
    assert texts and texts[0].startswith("Thương mại Việt Nam") and not any(t.startswith("Tuy nhiên") for t in texts)
    assert editorial.related("Việt Nam xuất siêu sang EU 36 tỷ USD trong 9 tháng",
                             "Thương mại Việt Nam EU 9 tháng, xuất siêu sang EU đạt 36 tỷ USD")
    assert not editorial.related("Việt Nam xuất siêu sang EU 36 tỷ USD",
                                 "Bộ trưởng Bộ Tài chính ký Hiệp định tránh đánh thuế hai lần với Liechtenstein")


def test_gold_is_one_world_sentence_without_silver():
    s = Snapshot.model_validate_json((ROOT / "tests/fixtures/snapshot.json").read_text(encoding="utf-8"))
    body = ("Giá vàng hôm nay\n\nGiá vàng quốc tế dao động quanh 4.179 USD/ounce sau khi dữ liệu lạm phát tháng 9 của Mỹ "
            "cao hơn dự báo. Giá bạc thế giới tăng 0,9% lên 57,5 USD/ounce trong phiên giao dịch sáng nay.")
    s.sources["news_gold"] = Source(id="news_gold", url="https://vneconomy.vn/g.htm", published_at=s.as_of - timedelta(hours=1),
                                    retrieved_at=s.as_of, text=body, kind="news", text_scope="article")
    s.observations["GOLD"].daily_pct = Decimal("0.3")
    result, audit = editorial.domestic_sections(s, {})
    gold = resolve(result["energy_metals"].paragraphs[1], s)
    assert gold.count(". ") == 0 and "bạc" not in gold and gold.startswith("Giá vàng quốc tế")


def test_page3_rows_and_disclaimer_blank_lines_match_sample(tmp_path):
    """Góp ý 20.10: cuối trang 3 trống vì dòng bảng giá hàng hóa và dòng trống giữa các đoạn miễn trừ bị ép còn 1pt."""
    doc = _render(tmp_path)
    cell = doc.tables[2].rows[1].cells[0]
    table = cell.tables[0]
    heights = [tr.trPr.find(qn("w:trHeight")).get(qn("w:val")) for tr in table._tbl.findall(qn("w:tr"))]
    assert heights[0] == "390" and set(heights[1:-1]) == {"195"} and heights[-1] == "255"
    disclaimer = doc.tables[2].rows[3].cells[0]
    blanks = [p for p in disclaimer.paragraphs if not p.text.strip()]
    assert len(blanks) >= 3
    for p in blanks:
        rule = p.paragraph_format.line_spacing
        assert rule == 1 or rule is None, rule        # dòng đơn cỡ 10, không còn 1pt
        assert p._p.pPr.find(qn("w:rPr")).find(qn("w:sz")).get(qn("w:val")) == "20"


def test_right_column_indent_matches_sample_and_china_is_short(tmp_path):
    from bttn.rendering import RIGHT_EDGE_INDENT
    assert RIGHT_EDGE_INDENT.pt == 8.1          # 162 dxa như tiêu đề và dòng "Dự kiến" của bản mẫu
    assert editorial.rules()["word_limits"]["china"] == [50, 62]
