"""Bản tin tuần: file sinh ra phải giữ nguyên định dạng của hai file mẫu."""
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from bttn.weekly import layout
from bttn.weekly.content import CalendarEvent, ChartData
from bttn.weekly.extract import extract_weekly
from bttn.weekly.render import render_weekly

ROOT = Path(__file__).resolve().parents[1]
TEMPLATES = {"vi": ROOT / "templates/weekly/BT_TTT_template_VIE.docx",
             "en": ROOT / "templates/weekly/BT_TTT_template_ENG.docx"}


def test_filename_and_labels_follow_sample():
    friday = date(2026, 10, 16)
    assert layout.weekly_filename(friday, "vi") == "BT_TTT2026.12.10-2026.16.10_VIE.docx"
    assert layout.weekly_filename(friday, "en") == "BT_TTT2026.12.10-2026.16.10_ENG.docx"
    assert layout.week_label(friday, "vi") == "Tuần từ 12.10 – 16.10.2026"
    assert layout.week_label(friday, "en") == "Week 42: 12.10 - 16.10.2026"
    assert layout.calendar_title(date(2026, 10, 19), "en") == "Economic calendar (Week 43: 19.10.2026 – 23.10.2026)"


@pytest.mark.parametrize("lang", ["vi", "en"])
def test_round_trip_keeps_content_and_static_parts(tmp_path, lang):
    content = extract_weekly(TEMPLATES[lang], lang, image_dir=tmp_path / "img")
    out = render_weekly(content, TEMPLATES[lang], tmp_path / "out.docx", split=False)
    again = extract_weekly(out, lang)
    a, b = content.model_dump(exclude={"images"}), again.model_dump(exclude={"images"})
    assert a == b
    # Phần cố định (đầu trang, miễn trừ, liên hệ) giống hệt mẫu, kể cả định dạng.
    tpl, new = Document(TEMPLATES[lang]).tables[0], Document(out).tables[0]
    for r in (0, 1, 14, 15):
        assert tpl.rows[r]._tr.xml == new.rows[r]._tr.xml


@pytest.mark.parametrize("lang", ["vi", "en"])
def test_three_pages_with_template_geometry(tmp_path, lang):
    content = extract_weekly(TEMPLATES[lang], lang)
    out = render_weekly(content, TEMPLATES[lang], tmp_path / "out.docx")
    doc = Document(out)
    grid = lambda t: [g.get(qn("w:w")) for g in t._tbl.tblGrid.findall(qn("w:gridCol"))]  # noqa: E731
    tpl = Document(TEMPLATES[lang]).tables[0]
    assert [len(t.rows) for t in doc.tables] == [len(p) for p in layout.PAGES]
    assert all(grid(t) == grid(tpl) for t in doc.tables)
    # Ngắt trang bằng "page break before" trên đoạn ngăn cách, không dùng ký tự ngắt trang
    # (ký tự ngắt trang làm Word sinh trang thứ 4 khi trang trước vừa đầy).
    assert not [br for br in doc.element.body.iter(qn("w:br")) if br.get(qn("w:type")) == "page"]
    assert len(list(doc.element.body.iter(qn("w:pageBreakBefore")))) == 2
    # Đoạn cuối văn bản thật thấp để không bị đẩy sang trang 4
    last = doc.element.body.findall(qn("w:p"))[-1]
    assert last.find(qn("w:pPr")).find(qn("w:spacing")).get(qn("w:lineRule")) == "exact"
    # 4 biểu đồ trang 1 không cao quá giới hạn
    wp = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"
    page1 = doc.tables[0]
    heights = [int(e.get("cy")) for e in page1._tbl.iter(wp) if int(e.get("cx")) > 3000000]
    assert len(heights) == 4 and max(heights) <= layout.PAGE1_CHART_MAX_CM * 360000
    page2 = doc.tables[1]
    img = [int(e.get("cy")) for e in page2._tbl.iter(wp)]
    assert len(img) == 6 and max(img) <= max(layout.PAGE2_IMAGE_MAX_CM.values()) * 360000


def test_new_content_charts_images_and_calendar(tmp_path):
    lang = "vi"
    content = extract_weekly(TEMPLATES[lang], lang)
    content.news = ["Tin một.", "Tin hai.", "Tin ba.", "Tin bốn."]
    content.sections["interbank"].body = ["Đoạn mới thứ nhất.", "Đoạn mới thứ hai."]
    content.sections["interbank"].forecast = "Dự kiến: thử nghiệm."
    content.charts["credit"] = ChartData(categories=["T1/2026", "T2/2026"], series=[[0.01, 0.02], [0.03, None]])
    png = tmp_path / "x.png"
    Image.new("RGB", (40, 20), "red").save(png)
    content.images = {"eurusd": str(png)}
    content.calendar.events = [CalendarEvent(date="19.10.2026", event="A", country="Mỹ", bold=True),
                               CalendarEvent(event="B", country="Nhật Bản", forecast="1%", previous="2%")]
    out = render_weekly(content, TEMPLATES[lang], tmp_path / "out.docx", split=False)
    again = extract_weekly(out, lang, image_dir=tmp_path / "img2")
    assert again.news == content.news
    assert again.sections["interbank"].body == content.sections["interbank"].body
    assert again.sections["interbank"].forecast == "Dự kiến: thử nghiệm."
    assert again.charts["credit"].categories == ["T1/2026", "T2/2026"]
    assert again.charts["credit"].series == [[0.01, 0.02], [0.03, None]]
    assert Path(again.images["eurusd"]).read_bytes() == png.read_bytes()
    assert [(e.event, e.bold) for e in again.calendar.events] == [("A", True), ("B", False)]
    # Đoạn dự báo vẫn đậm, màu đỏ như mẫu
    doc = Document(out)
    tc = doc.tables[0].rows[3]._tr.tc_lst[0]
    fc = [p for p in tc.findall(qn("w:p")) if "thử nghiệm" in "".join(t.text or "" for t in p.iter(qn("w:t")))][0]
    rpr = fc.find(qn("w:r")).find(qn("w:rPr"))
    assert rpr.find(qn("w:color")).get(qn("w:val")) == "C00000" and rpr.find(qn("w:b")) is not None


def test_chart_series_count_must_match_template(tmp_path):
    content = extract_weekly(TEMPLATES["vi"], "vi")
    content.charts["interbank"] = ChartData(categories=["a"], series=[[1.0]])
    with pytest.raises(ValueError):
        render_weekly(content, TEMPLATES["vi"], tmp_path / "x.docx")
